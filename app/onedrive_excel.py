"""Espelha o roteiro do dia (aba "Roteiro por Técnico") numa planilha Excel Online
já existente, pra o time operacional acompanhar sem precisar copiar isso à mão
todo dia. O arquivo real ("Roteiro_2026") não é dono da conta que faz login aqui
— é de outra pessoa e só foi compartilhado ("Compartilhado comigo" no OneDrive) —
por isso ele é localizado via `/me/drive/sharedWithMe` (por nome), não por um
caminho dentro do próprio drive.

Cada dia ganha uma aba nova dentro do MESMO arquivo (histórico "quando tal OS saiu
com qual técnico") — abas de dias anteriores nunca são tocadas; só a aba do dia
corrente é recriada/limpa se a roteirização for reprocessada no mesmo dia.

Autenticação via MSAL (fluxo de device code — a conta é pessoal, não Microsoft 365
corporativo, então client-credentials/app-only não se aplica). O token fica em cache
localmente (`onedrive_profile/token_cache.bin`, incluído em PASTAS_WORKSPACE pra ser
arquivado/restaurado por conta, mesmo padrão de whatsapp_profile/) e nunca é enviado
pro backend — é um segredo mais sensível que uma senha de portal, e o backend não
criptografa o campo de config genérico (só as credenciais Mobyan/OGEA são).

A automação nunca dispara o login interativo sozinha: `enviar_roteiro_dia` só tenta
aquisição silenciosa de token. Se não houver sessão em cache, levanta
`PrecisaLoginError` e quem chamou decide o que fazer (no caso, `gerar_roteirizacao.py`
só avisa no log e segue sem travar a planilha local). O login interativo (device
code) só roda via `app/conectar_onedrive.py`, executado manualmente pelo usuário.
"""

import os
import webbrowser

import msal
import requests
from openpyxl.utils import get_column_letter

from caminho_base import BASE_DIR

PASTA_PERFIL_ONEDRIVE = BASE_DIR / "onedrive_profile"
ARQUIVO_CACHE_TOKEN = PASTA_PERFIL_ONEDRIVE / "token_cache.bin"

AUTORIDADE = "https://login.microsoftonline.com/consumers"
ESCOPOS = ["https://graph.microsoft.com/Files.ReadWrite"]
GRAFO_BASE = "https://graph.microsoft.com/v1.0"

MESES_ABREV = {
    1: "JAN", 2: "FEV", 3: "MAR", 4: "ABR", 5: "MAI", 6: "JUN",
    7: "JUL", 8: "AGO", 9: "SET", 10: "OUT", 11: "NOV", 12: "DEZ",
}
DIAS_SEMANA_ABREV = {0: "SEG", 1: "TER", 2: "QUA", 3: "QUI", 4: "SEX", 5: "SAB", 6: "DOM"}


class PrecisaLoginError(Exception):
    """Não há sessão OneDrive em cache — é preciso rodar app/conectar_onedrive.py."""


def esta_configurado():
    return bool(os.getenv("ONEDRIVE_CLIENT_ID", "").strip()) and bool(
        os.getenv("ONEDRIVE_ARQUIVO_CAMINHO", "").strip()
    )


def _client_id():
    client_id = os.getenv("ONEDRIVE_CLIENT_ID", "").strip()
    if not client_id:
        raise RuntimeError("ONEDRIVE_CLIENT_ID não configurado para esta conta.")
    return client_id


def _nome_arquivo():
    nome = os.getenv("ONEDRIVE_ARQUIVO_CAMINHO", "").strip()
    if not nome:
        raise RuntimeError("ONEDRIVE_ARQUIVO_CAMINHO não configurado para esta conta.")
    return nome


def _carregar_cache():
    cache = msal.SerializableTokenCache()
    if ARQUIVO_CACHE_TOKEN.exists():
        cache.deserialize(ARQUIVO_CACHE_TOKEN.read_text(encoding="utf-8"))
    return cache


def _salvar_cache(cache):
    if not cache.has_state_changed:
        return
    PASTA_PERFIL_ONEDRIVE.mkdir(parents=True, exist_ok=True)
    ARQUIVO_CACHE_TOKEN.write_text(cache.serialize(), encoding="utf-8")
    ARQUIVO_CACHE_TOKEN.chmod(0o600)


def _app_msal(cache):
    return msal.PublicClientApplication(
        client_id=_client_id(),
        authority=AUTORIDADE,
        token_cache=cache,
    )


def obter_token(interativo=False):
    """Tenta aquisição silenciosa (renovação automática via refresh token em
    cache). Só faz o fluxo de device code (que exige o usuário abrir o
    navegador) se `interativo=True` — usado só por app/conectar_onedrive.py."""
    cache = _carregar_cache()
    app = _app_msal(cache)

    contas = app.get_accounts()
    if contas:
        resultado = app.acquire_token_silent(ESCOPOS, account=contas[0])
        if resultado and "access_token" in resultado:
            _salvar_cache(cache)
            return resultado["access_token"]

    if not interativo:
        raise PrecisaLoginError(
            "Nenhuma sessão do OneDrive em cache. Rode 'Conectar OneDrive' "
            "(app/conectar_onedrive.py) uma vez pra autorizar o acesso."
        )

    fluxo = app.initiate_device_flow(scopes=ESCOPOS)
    if "user_code" not in fluxo:
        raise RuntimeError(
            f"Não consegui iniciar o login do OneDrive: {fluxo.get('error_description', fluxo)}"
        )

    print(f"Abra {fluxo['verification_uri']} e digite o código: {fluxo['user_code']}")
    try:
        webbrowser.open(fluxo["verification_uri"])
    except Exception:
        pass

    resultado = app.acquire_token_by_device_flow(fluxo)
    if "access_token" not in resultado:
        raise RuntimeError(
            f"Login do OneDrive não concluído: {resultado.get('error_description', resultado)}"
        )

    _salvar_cache(cache)
    return resultado["access_token"]


def nome_aba_do_dia(data):
    return f"{data.day:02d} {MESES_ABREV[data.month]} {DIAS_SEMANA_ABREV[data.weekday()]}"


def _chamar_grafo(token, metodo, url, **kwargs):
    cabecalhos = kwargs.pop("headers", {})
    cabecalhos["Authorization"] = f"Bearer {token}"
    resposta = requests.request(metodo, url, headers=cabecalhos, timeout=30, **kwargs)
    if not resposta.ok:
        raise RuntimeError(
            f"Chamada à Microsoft Graph falhou ({resposta.status_code}): {resposta.text[:500]}"
        )
    return resposta


def _resolver_item_arquivo(token, nome_arquivo):
    """Localiza um arquivo que foi compartilhado com a conta logada (não é dela)
    procurando por nome em 'Compartilhado comigo'. Retorna (drive_id, item_id) —
    diferente de um arquivo próprio, um item compartilhado vive no drive de quem
    o compartilhou, então toda chamada seguinte precisa endereçar
    /drives/{drive_id}/items/{item_id}, não /me/drive/items/{item_id}."""
    url = f"{GRAFO_BASE}/me/drive/sharedWithMe"
    itens = _chamar_grafo(token, "GET", url).json().get("value", [])

    alvo = nome_arquivo.strip().lower()
    for item in itens:
        remoto = item.get("remoteItem") or item
        nome_item = (remoto.get("name") or item.get("name") or "").lower()
        if nome_item == alvo or nome_item.startswith(alvo + "."):
            referencia = remoto.get("parentReference") or {}
            drive_id = referencia.get("driveId")
            item_id = remoto.get("id")
            if drive_id and item_id:
                return drive_id, item_id

    raise RuntimeError(
        f"Não encontrei '{nome_arquivo}' em 'Compartilhado comigo' no OneDrive. "
        "Confirme o nome exato do arquivo e se ele ainda está compartilhado com esta conta."
    )


def verificar_conexao(token):
    """Confirma que o arquivo configurado em ONEDRIVE_ARQUIVO_CAMINHO é
    alcançável com o token informado. Retorna (drive_id, item_id)."""
    return _resolver_item_arquivo(token, _nome_arquivo())


def _base_item(drive_id, item_id):
    return f"{GRAFO_BASE}/drives/{drive_id}/items/{item_id}"


def _localizar_ou_criar_aba(token, drive_id, item_id, nome):
    url_worksheets = f"{_base_item(drive_id, item_id)}/workbook/worksheets"
    abas = _chamar_grafo(token, "GET", url_worksheets).json().get("value", [])

    for aba in abas:
        if aba.get("name") == nome:
            return aba["id"]

    resposta = _chamar_grafo(
        token, "POST", f"{url_worksheets}/add", json={"name": nome}
    )
    return resposta.json()["id"]


def _url_range(drive_id, item_id, ws_id, endereco):
    return (
        f"{_base_item(drive_id, item_id)}/workbook/worksheets/{ws_id}"
        f"/range(address='{endereco}')"
    )


def _limpar_aba(token, drive_id, item_id, ws_id):
    url = f"{_url_range(drive_id, item_id, ws_id, 'A1:AZ1000')}/clear"
    _chamar_grafo(token, "POST", url, json={"applyTo": "All"})


def _escrever_valores(token, drive_id, item_id, ws_id, endereco, valores):
    _chamar_grafo(
        token, "PATCH", _url_range(drive_id, item_id, ws_id, endereco), json={"values": valores}
    )


def _formatar_range(token, drive_id, item_id, ws_id, endereco, **propriedades_fonte):
    url = f"{_url_range(drive_id, item_id, ws_id, endereco)}/format/font"
    _chamar_grafo(token, "PATCH", url, json=propriedades_fonte)


def enviar_roteiro_dia(matriz, origens_colunas, tecnicos, hoje):
    """`matriz`/`origens_colunas`/`tecnicos` são exatamente os já calculados por
    gerar_roteirizacao.py::gerar_planilha() pra aba local "Roteiro por Técnico" —
    nada é recalculado aqui. Levanta exceção em qualquer falha; quem chama decide
    se isso deve travar o fluxo (não deve, no caso de gerar_roteirizacao.py)."""
    token = obter_token(interativo=False)
    drive_id, item_id = _resolver_item_arquivo(token, _nome_arquivo())
    ws_id = _localizar_ou_criar_aba(token, drive_id, item_id, nome_aba_do_dia(hoje))

    _limpar_aba(token, drive_id, item_id, ws_id)

    max_linhas = max((len(matriz[tecnico]) for tecnico in tecnicos), default=0)
    ultima_coluna = get_column_letter(max(len(tecnicos), 1))

    _escrever_valores(token, drive_id, item_id, ws_id, f"A1:{ultima_coluna}1", [tecnicos])
    _formatar_range(token, drive_id, item_id, ws_id, f"A1:{ultima_coluna}1", bold=True)

    if max_linhas == 0:
        return

    linhas = []
    for indice_linha in range(max_linhas):
        linha = []
        for tecnico in tecnicos:
            valores_tecnico = matriz[tecnico]
            linha.append(valores_tecnico[indice_linha] if indice_linha < len(valores_tecnico) else None)
        linhas.append(linha)

    _escrever_valores(token, drive_id, item_id, ws_id, f"A2:{ultima_coluna}{max_linhas + 1}", linhas)

    for indice_coluna, tecnico in enumerate(tecnicos, start=1):
        letra = get_column_letter(indice_coluna)
        origens = origens_colunas.get(tecnico, [])
        qtd_mobyan = sum(1 for origem in origens if origem == "MOBYAN")
        qtd_ogea = sum(1 for origem in origens if origem == "OGEA")

        if qtd_mobyan:
            _formatar_range(
                token, drive_id, item_id, ws_id, f"{letra}2:{letra}{qtd_mobyan + 1}", color="#FF0000"
            )
        if qtd_ogea:
            inicio = qtd_mobyan + 2
            fim = qtd_mobyan + qtd_ogea + 1
            _formatar_range(
                token, drive_id, item_id, ws_id, f"{letra}{inicio}:{letra}{fim}", color="#0070C0"
            )
